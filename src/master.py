"""JPX の ETF データベース xlsx を取得・パースする（銘柄マスタ）。

xlsx の href はページ依存のため、JPX の投資家向けページから動的に抽出する。
取得失敗時はローカルに残っている xlsx をフォールバックとして使う。
"""

from __future__ import annotations

import logging
import re
from datetime import date, datetime
from pathlib import Path
from urllib.parse import urljoin

import requests
from openpyxl import load_workbook

log = logging.getLogger("master")

ROOT = Path(__file__).resolve().parents[1]
JPX_PAGE_URL = "https://www.jpx.co.jp/equities/products/etfs/investors/index.html"
FALLBACK_XLSX = ROOT / "nlsgeu000000vx9t.xlsx"
DOWNLOADED_XLSX = ROOT / "etf_database.xlsx"

SHEET_NAME = "データベース"
HEADER_ROW = 4  # 1 行目タイトル、4 行目がヘッダ（A 列は空白列、B 列「コード」始まり）
DATA_START_ROW = 5

# 詳細ページのセクション分け（列名 → セクション）
SECTIONS: list[tuple[str, list[str]]] = [
    ("基本情報", ["コード", "名称", "略称", "連動指標名", "カテゴリー", "管理会社", "上場日"]),
    ("売買・流動性", ["売買単位(口)", "立会内", "ToSTNeT", "OTC※日本証券業協会公表値を基に算出"]),
    ("規模・コスト", ["純資産総額（2025年末・百万円）", "総経費率（年率・税込）", "信託報酬（年率・税抜）", "信託財産留保額（基準価額比の●%）"]),
    ("分配・価格", ["分配金支払基準日", "株価（2025年末・円）", "実績分配金（2025年末・円）", "実績分配金利回り（2025年末）"]),
    ("iNAV・PCF", ["iNAV配信", "PCF公表方法", "PCF掲載URL(日本語)", "PCF掲載URL(英語)", "PCF更新頻度"]),
    ("制度・分類", ["根拠法の分類", "税法上の分類", "英語情報の掲載URL", "ブラックアウト期間のカレンダー掲載URL（日本語)", "ブラックアウト期間のカレンダー掲載URL（英語)", "特記事項", "最終更新日"]),
    ("設定・交換スケジュール", [
        "設定スケジュール【ETF設定に係る債務引受対象でない場合】",
        "交換（解約）スケジュール【ETF設定に係る債務引受対象でない場合】",
        "設定スケジュール【ETF設定に係る債務引受対象である場合】",
        "交換（解約）スケジュール【ETF設定に係る債務引受対象である場合】",
        "設定単位①", "設定単位②", "設定単位③",
        "交換単位①", "交換単位②", "交換単位③",
        "1口あたりの基準価額（2025年末・円）",
        "設定概算金額（百万円）", "交換（解約）概算金額（百万円）",
        "CredNex", "設定", "交換",
    ]),
]


def normalize_header(v: object) -> str:
    return str(v).replace("\n", "").strip()


def _normalize_value(v: object) -> object:
    if v is None:
        return None
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def fetch_xlsx(timeout: int = 30) -> Path | None:
    """JPX ページから .xlsx の href を抽出してダウンロード。失敗時は None。"""
    try:
        res = requests.get(JPX_PAGE_URL, timeout=timeout)
        res.raise_for_status()
        m = re.search(r'href="([^"]+\.xlsx)"', res.text)
        if not m:
            log.warning("xlsx の href が見つかりません")
            return None
        url = urljoin(JPX_PAGE_URL, m.group(1))
        r = requests.get(url, timeout=timeout)
        r.raise_for_status()
        DOWNLOADED_XLSX.write_bytes(r.content)
        log.info("xlsx をダウンロード: %s", url)
        return DOWNLOADED_XLSX
    except Exception as e:
        log.warning("xlsx ダウンロード失敗（フォールバックを使用）: %s", e)
        return None


def xlsx_path(prefer_download: bool = True) -> Path:
    if prefer_download:
        p = fetch_xlsx()
        if p is not None:
            return p
    if FALLBACK_XLSX.is_file():
        return FALLBACK_XLSX
    if DOWNLOADED_XLSX.is_file():
        return DOWNLOADED_XLSX
    raise FileNotFoundError(f"xlsx が見つかりません: {FALLBACK_XLSX} / {DOWNLOADED_XLSX}")


def load_master(path: Path | None = None) -> list[dict]:
    """xlsx を読み、銘柄レコードのリストを返す。各レコードは code/name/category/fields。"""
    path = path or xlsx_path()
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[SHEET_NAME]
    rows = ws.iter_rows(values_only=True)
    for _ in range(HEADER_ROW - 1):
        next(rows)
    headers = [normalize_header(h) for h in next(rows)]

    records: list[dict] = []
    for row in rows:
        code = row[1] if len(row) > 1 else None
        if code is None:
            continue
        fields: dict[str, object] = {}
        for i in range(1, len(headers)):
            if headers[i]:
                fields[headers[i]] = _normalize_value(row[i] if i < len(row) else None)
        records.append({
            "code": str(code),
            "name": str(fields.get("名称") or ""),
            "category": str(fields.get("カテゴリー") or ""),
            "fields": fields,
        })
    wb.close()
    return records
